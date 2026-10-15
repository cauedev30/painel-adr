"""Grava os 6 brutos e o gabarito em disco, cada arquivo do jeito que o sistema de origem exporta.

CRM: CSV em UTF-8 com BOM e `;` (D2.2). ERP: CSV em Latin-1 e `;` (D2.3). Planilhas: .xlsx com uma aba
por chave. Todo arquivo de texto grava fim de linha `\\n` explícito: o pandas no Windows grava `\\r\\n`.
"""

import json
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter

from etl.gerar import sujeira

CODIFICACAO_CSV = {sujeira.CRM: "utf-8-sig", sujeira.ERP_PARCELAS: "latin-1", sujeira.ERP_BAIXAS: "latin-1"}
LARGURA_MAXIMA = 50


FORMATO_DATA = "DD/MM/YYYY"
FORMATO_DATA_HORA = "DD/MM/YYYY HH:MM"
FORMATO_PONTOS_PERCENTUAIS = "0.00"  # a coluna "(%)" já vem em pontos: 3,26 é 3,26%


def _formato(serie: pd.Series, titulo: str) -> str | None:
    """Formato de exibição da coluna no Excel; None mantém o padrão."""
    if pd.api.types.is_datetime64_any_dtype(serie):
        com_hora = (serie.dropna().dt.floor("D") != serie.dropna()).any()
        return FORMATO_DATA_HORA if com_hora else FORMATO_DATA
    if titulo.endswith("(%)"):
        return FORMATO_PONTOS_PERCENTUAIS
    return None


def _largura(serie: pd.Series, titulo: str, formato: str | None) -> int:
    """Largura da coluna no Excel: o maior texto da coluna, com folga."""
    if formato in (FORMATO_DATA, FORMATO_DATA_HORA):
        maior = len(formato)
    else:
        maior = int(serie.dropna().astype(str).str.len().max()) if serie.notna().any() else 0
    return min(max(maior, len(titulo)) + 2, LARGURA_MAXIMA)


def _gravar_xlsx(abas: dict[str, pd.DataFrame], caminho: Path) -> None:
    # O formato vai célula a célula: o pandas 3 ignora o date_format e o datetime_format do ExcelWriter
    with pd.ExcelWriter(caminho, engine="openpyxl") as escritor:
        for aba, df in abas.items():
            df.to_excel(escritor, sheet_name=aba, index=False)
            planilha = escritor.sheets[aba]
            for n, coluna in enumerate(df.columns, start=1):
                letra = get_column_letter(n)
                formato = _formato(df[coluna], coluna)
                if formato:
                    for (celula,) in planilha.iter_rows(min_row=2, min_col=n, max_col=n):
                        celula.number_format = formato
                planilha.column_dimensions[letra].width = _largura(df[coluna], coluna, formato)


def gravar(brutos: dict, gabarito: dict, pasta_brutos: Path, arq_gabarito: Path) -> dict[str, int]:
    """Grava os brutos e o gabarito; devolve as linhas gravadas por arquivo."""
    pasta_brutos.mkdir(parents=True, exist_ok=True)
    linhas = {}
    for arquivo, conteudo in brutos.items():
        caminho = pasta_brutos / arquivo
        if arquivo in CODIFICACAO_CSV:
            conteudo.to_csv(caminho, sep=";", encoding=CODIFICACAO_CSV[arquivo], lineterminator="\n", index=False)
            linhas[arquivo] = len(conteudo)
        else:
            _gravar_xlsx(conteudo, caminho)
            linhas[arquivo] = sum(len(df) for df in conteudo.values())

    arq_gabarito.parent.mkdir(parents=True, exist_ok=True)
    with open(arq_gabarito, "w", encoding="utf-8", newline="\n") as f:
        f.write(json.dumps(gabarito, ensure_ascii=False, indent=2) + "\n")
    return linhas
